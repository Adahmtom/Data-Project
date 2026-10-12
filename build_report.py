"""
Monthly KPI report for FormWise (made-up company, simulated data).

Pulls metrics from the database with SQL and writes an Excel workbook where every
KPI is a formula driven by one input cell, the report month. Pick a different month
from the dropdown and the whole workbook updates. Also writes a short text summary
of the month.

Usage:
    python build_report.py --month 2026-06
"""
import argparse, os, sqlite3, subprocess, sys
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
DB = os.path.join(HERE, "data", "formwise.db")


def connect():
    if not os.path.exists(DB):
        subprocess.run([sys.executable, os.path.join(HERE, "data", "generate_data.py")], check=True)
    return sqlite3.connect(DB)


def q(sql, con):
    return pd.read_sql_query(sql, con)


FONT = "Arial"
HEAD = PatternFill("solid", fgColor="1F3864")
INPUT = PatternFill("solid", fgColor="FFFF00")
BAND = PatternFill("solid", fgColor="EEF2F7")
THIN = Border(bottom=Side(style="thin", color="BBBBBB"))


def extract(con):
    base = q("""
    SELECT strftime('%Y-%m', v.first_visit_ts) AS month, v.channel, v.device,
           COUNT(*) AS visitors, COUNT(s.user_id) AS purchases, COALESCE(SUM(p.amount), 0) AS new_revenue
    FROM visitors v
    LEFT JOIN subscriptions s ON s.visitor_id = v.visitor_id
    LEFT JOIN payments p ON p.user_id = s.user_id AND p.payment_type = 'new'
    GROUP BY month, v.channel, v.device ORDER BY month, v.channel, v.device""", con)
    renew = q("""SELECT strftime('%Y-%m', payment_ts) AS month, SUM(amount) AS renewal_revenue,
                        COUNT(*) AS renewals
                 FROM payments WHERE payment_type = 'renewal' GROUP BY month""", con)
    active = q("""
    WITH months AS (SELECT DISTINCT strftime('%Y-%m', first_visit_ts) AS month FROM visitors)
    SELECT m.month,
           (SELECT COUNT(*) FROM subscriptions s
             WHERE s.start_ts < date(m.month || '-01', '+1 month')
               AND (s.end_ts IS NULL OR s.end_ts >= date(m.month || '-01', '+1 month'))) AS active_subscribers_eom
    FROM months m ORDER BY m.month""", con)
    dq = q("""
    WITH ev AS (
      SELECT strftime('%Y-%m', event_ts) AS month, COUNT(*) AS purchase_events,
             SUM(user_id IS NULL) AS missing_user_id,
             COUNT(order_id) - COUNT(DISTINCT order_id) AS duplicate_events
      FROM events WHERE event_name = 'purchase' GROUP BY month),
    bill AS (SELECT strftime('%Y-%m', payment_ts) AS month, COUNT(*) AS billed_new_subs
             FROM payments WHERE payment_type = 'new' GROUP BY month)
    SELECT b.month, b.billed_new_subs, e.purchase_events, e.missing_user_id, e.duplicate_events
    FROM bill b JOIN ev e ON e.month = b.month ORDER BY b.month""", con)
    months = sorted(base.month.unique())
    renew = pd.DataFrame({"month": months}).merge(renew, on="month", how="left").fillna(0)
    dq = dq[dq.month.isin(months)]
    return base, renew, active[active.month.isin(months)], dq, months


def style_header(ws, row, ncols, start_col=1):
    for c in range(start_col, start_col + ncols):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF"); cell.fill = HEAD
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def set_fonts(ws):
    for row in ws.iter_rows():
        for c in row:
            if c.font and c.font.name != FONT:
                c.font = Font(name=FONT, bold=c.font.bold, color=c.font.color, size=c.font.size, italic=c.font.italic)


def build(month, out_dir):
    con = connect()
    base, renew, active, dq, months = extract(con)
    if month not in months:
        raise SystemExit(f"Month {month} not in data ({months[0]} to {months[-1]})")
    wb = Workbook()

    # Data (raw extract)
    wsd = wb.active; wsd.title = "Data"
    cols = ["month", "channel", "device", "visitors", "purchases", "new_revenue"]
    wsd.append(cols); style_header(wsd, 1, len(cols))
    for r in base[cols].itertuples(index=False):
        wsd.append(list(r))
    n_data = len(base) + 1
    for c, w in zip("ABCDEF", [10, 16, 10, 11, 11, 13]):
        wsd.column_dimensions[c].width = w
    for r in range(2, n_data + 1):
        wsd.cell(row=r, column=6).number_format = "$#,##0"
    wsd.freeze_panes = "A2"
    D = lambda col: f"Data!${col}$2:${col}${n_data}"

    # Monthly
    wsm = wb.create_sheet("Monthly")
    mh = ["Month", "Visitors", "Purchases", "Conversion", "New revenue", "Renewal revenue", "Total revenue",
          "Renewals", "Active subscribers (end of month)"]
    wsm.append(mh); style_header(wsm, 1, len(mh))
    for i, m in enumerate(months, start=2):
        rr = renew.set_index("month").loc[m]
        wsm.append([m, f"=SUMIFS({D('D')},{D('A')},A{i})", f"=SUMIFS({D('E')},{D('A')},A{i})",
                    f"=IFERROR(C{i}/B{i},0)", f"=SUMIFS({D('F')},{D('A')},A{i})",
                    float(rr.renewal_revenue), f"=E{i}+F{i}", int(rr.renewals),
                    int(active.set_index("month").loc[m, "active_subscribers_eom"])])
        for c, fmt in zip("BCDEFGHI", ["#,##0", "#,##0", "0.00%", "$#,##0", "$#,##0", "$#,##0", "#,##0", "#,##0"]):
            wsm[f"{c}{i}"].number_format = fmt
        if i % 2 == 0:
            for c in range(1, 10): wsm.cell(row=i, column=c).fill = BAND
    last = len(months) + 1
    wsm[f"A{last+2}"] = "Renewal revenue, renewals and active subscribers are loaded from the billing system by build_report.py."
    wsm[f"A{last+2}"].font = Font(name=FONT, italic=True, size=9, color="666666")
    for c, w in zip("ABCDEFGHI", [10, 11, 11, 12, 13, 15, 14, 11, 18]):
        wsm.column_dimensions[c].width = w
    wsm.freeze_panes = "B2"

    ch = LineChart(); ch.title = "Visit to purchase conversion"; ch.height, ch.width = 7.5, 17
    ch.add_data(Reference(wsm, min_col=4, min_row=1, max_row=last), titles_from_data=True)
    ch.set_categories(Reference(wsm, min_col=1, min_row=2, max_row=last)); ch.y_axis.numFmt = "0.0%"
    ch.legend = None; wsm.add_chart(ch, "K2")
    bc = BarChart(); bc.type = "col"; bc.grouping = "stacked"; bc.overlap = 100
    bc.title = "Revenue by month (new + renewal)"; bc.height, bc.width = 7.5, 17
    bc.add_data(Reference(wsm, min_col=5, max_col=6, min_row=1, max_row=last), titles_from_data=True)
    bc.set_categories(Reference(wsm, min_col=1, min_row=2, max_row=last)); bc.y_axis.numFmt = "$#,##0"
    wsm.add_chart(bc, "K18")

    # KPI Summary
    ws = wb.create_sheet("KPI Summary", 0)
    ws["A1"] = "FormWise Monthly KPI Report"; ws["A1"].font = Font(name=FONT, bold=True, size=16, color="1F3864")
    ws["A2"] = "Made-up company, simulated data. Change the yellow cell to see another month."
    ws["A2"].font = Font(name=FONT, italic=True, size=9, color="666666")
    ws["A4"] = "Report month"; ws["A4"].font = Font(name=FONT, bold=True)
    ws["B4"] = month; ws["B4"].fill = INPUT; ws["B4"].font = Font(name=FONT, bold=True, color="0000FF")
    dv = DataValidation(type="list", formula1=f"=Monthly!$A$2:$A${last}", allow_blank=False)
    ws.add_data_validation(dv); dv.add("B4")
    ws["C4"] = "pick a month from the list"; ws["C4"].font = Font(name=FONT, italic=True, color="666666")
    ws["A5"] = "Row in Monthly"; ws["B5"] = f"=MATCH($B$4,Monthly!$A$2:$A${last},0)"
    ws["A5"].font = ws["B5"].font = Font(name=FONT, size=8, color="999999")

    hdr = ["KPI", "This month", "Prior month", "Change vs prior", "Same month last year", "Change vs last year"]
    for j, h in enumerate(hdr, 1): ws.cell(row=7, column=j, value=h)
    style_header(ws, 7, len(hdr))
    kpis = [("Visitors", "B", "#,##0"), ("Purchases (new subscriptions)", "C", "#,##0"),
            ("Visit to purchase conversion", "D", "0.00%"), ("New subscription revenue", "E", "$#,##0"),
            ("Renewal revenue", "F", "$#,##0"), ("Total revenue", "G", "$#,##0"),
            ("Active subscribers (end of month)", "I", "#,##0")]
    rng_ = lambda col: f"Monthly!${col}$2:${col}${last}"
    for i, (label, col, fmt) in enumerate(kpis, start=8):
        ws[f"A{i}"] = label
        ws[f"B{i}"] = f"=INDEX({rng_(col)},$B$5)"
        ws[f"C{i}"] = f'=IF($B$5>1,INDEX({rng_(col)},$B$5-1),"n/a")'
        ws[f"E{i}"] = f'=IF($B$5>12,INDEX({rng_(col)},$B$5-12),"n/a")'
        if fmt == "0.00%":   # rates: show change in percentage points
            ws[f"D{i}"] = f'=IF(ISNUMBER(C{i}),(B{i}-C{i})*100,"n/a")'
            ws[f"F{i}"] = f'=IF(ISNUMBER(E{i}),(B{i}-E{i})*100,"n/a")'
            ws[f"D{i}"].number_format = ws[f"F{i}"].number_format = '+0.00" pts";-0.00" pts"'
        else:
            ws[f"D{i}"] = f'=IF(ISNUMBER(C{i}),IFERROR(B{i}/C{i}-1,0),"n/a")'
            ws[f"F{i}"] = f'=IF(ISNUMBER(E{i}),IFERROR(B{i}/E{i}-1,0),"n/a")'
            ws[f"D{i}"].number_format = ws[f"F{i}"].number_format = "+0.0%;-0.0%"
        for c in "BCE": ws[f"{c}{i}"].number_format = fmt
        for c in "ABCDEF": ws[f"{c}{i}"].border = THIN
    kend = 7 + len(kpis)
    green, red = Font(name=FONT, color="2E7D32", bold=True), Font(name=FONT, color="C62828", bold=True)
    for col in "DF":
        rngc = f"{col}8:{col}{kend}"
        ws.conditional_formatting.add(rngc, FormulaRule(formula=[f"AND(ISNUMBER({col}8),{col}8>0)"], font=green))
        ws.conditional_formatting.add(rngc, FormulaRule(formula=[f"AND(ISNUMBER({col}8),{col}8<0)"], font=red))

    # Channel x device grid for the selected month (pivot-style with SUMIFS)
    r0 = kend + 3
    ws[f"A{r0-1}"] = "Conversion by channel and device (selected month)"; ws[f"A{r0-1}"].font = Font(name=FONT, bold=True, size=12)
    devices = ["desktop", "mobile", "tablet"]
    channels = sorted(base.channel.unique())
    heads = ["Channel"] + [f"{d} conv." for d in devices] + ["Visitors", "Purchases", "All devices conv.", "New revenue"]
    for j, h in enumerate(heads, 1): ws.cell(row=r0, column=j, value=h)
    style_header(ws, r0, len(heads))
    for k, chn in enumerate(channels, start=r0 + 1):
        ws[f"A{k}"] = chn
        for j, dev in enumerate(devices):
            col = get_column_letter(2 + j)
            ws[f"{col}{k}"] = (f'=IFERROR(SUMIFS({D("E")},{D("A")},$B$4,{D("B")},$A{k},{D("C")},"{dev}")'
                               f'/SUMIFS({D("D")},{D("A")},$B$4,{D("B")},$A{k},{D("C")},"{dev}"),0)')
            ws[f"{col}{k}"].number_format = "0.0%"
        ws[f"E{k}"] = f"=SUMIFS({D('D')},{D('A')},$B$4,{D('B')},$A{k})"
        ws[f"F{k}"] = f"=SUMIFS({D('E')},{D('A')},$B$4,{D('B')},$A{k})"
        ws[f"G{k}"] = f"=IFERROR(F{k}/E{k},0)"
        ws[f"H{k}"] = f"=SUMIFS({D('F')},{D('A')},$B$4,{D('B')},$A{k})"
        ws[f"E{k}"].number_format = ws[f"F{k}"].number_format = "#,##0"
        ws[f"G{k}"].number_format = "0.0%"; ws[f"H{k}"].number_format = "$#,##0"
    tot = r0 + 1 + len(channels)
    ws[f"A{tot}"] = "Total"; ws[f"A{tot}"].font = Font(name=FONT, bold=True)
    for c in "EFH":
        ws[f"{c}{tot}"] = f"=SUM({c}{r0+1}:{c}{tot-1})"; ws[f"{c}{tot}"].font = Font(name=FONT, bold=True)
    ws[f"G{tot}"] = f"=IFERROR(F{tot}/E{tot},0)"; ws[f"G{tot}"].font = Font(name=FONT, bold=True)
    ws[f"E{tot}"].number_format = ws[f"F{tot}"].number_format = "#,##0"
    ws[f"G{tot}"].number_format = "0.0%"; ws[f"H{tot}"].number_format = "$#,##0"
    ws[f"A{tot+1}"] = "Check: total visitors match KPI table"
    ws[f"E{tot+1}"] = f'=IF(E{tot}=B8,"OK","MISMATCH")'
    ws[f"A{tot+1}"].font = ws[f"E{tot+1}"].font = Font(name=FONT, italic=True, size=9, color="666666")

    cc = BarChart(); cc.type = "bar"; cc.title = "Conversion by channel (selected month)"
    cc.height, cc.width = 7.5, 14
    cc.add_data(Reference(ws, min_col=7, min_row=r0, max_row=tot - 1), titles_from_data=True)
    cc.set_categories(Reference(ws, min_col=1, min_row=r0 + 1, max_row=tot - 1))
    cc.x_axis.numFmt = "0%"; cc.legend = None
    ws.add_chart(cc, f"A{tot+3}")
    for c, w in zip("ABCDEFGH", [36, 14, 14, 16, 18, 16, 17, 14]):
        ws.column_dimensions[c].width = w
    ws.freeze_panes = "A8"

    # Data Quality
    wq = wb.create_sheet("Data Quality")
    qh = ["Month", "Billed new subs", "Purchase events", "Missing user_id", "Duplicate events",
          "Event vs billing gap", "Status"]
    wq.append(qh); style_header(wq, 1, len(qh))
    wq["I1"] = "Gap tolerance"; wq["J1"] = 0.03; wq["J1"].number_format = "0%"
    wq["I1"].font = Font(name=FONT, bold=True); wq["J1"].font = Font(name=FONT, color="0000FF"); wq["J1"].fill = INPUT
    wq["I2"] = "Status is CHECK when any user_id is missing or the gap exceeds the tolerance."
    wq["I2"].font = Font(name=FONT, italic=True, size=9, color="666666")
    for i, r in enumerate(dq.itertuples(index=False), start=2):
        wq.append([r.month, int(r.billed_new_subs), int(r.purchase_events), int(r.missing_user_id), int(r.duplicate_events)])
        wq[f"F{i}"] = f"=IFERROR(C{i}/B{i}-1,0)"; wq[f"F{i}"].number_format = "+0.0%;-0.0%"
        wq[f"G{i}"] = f'=IF(OR(D{i}>0,ABS(F{i})>$J$1),"CHECK","OK")'
    qlast = len(dq) + 1
    wq.conditional_formatting.add(f"G2:G{qlast}", FormulaRule(formula=['G2="CHECK"'], font=red,
                                  fill=PatternFill("solid", fgColor="FDECEA")))
    for c, w in zip("ABCDEFGHIJ", [10, 15, 15, 15, 15, 18, 10, 2, 16, 8]):
        wq.column_dimensions[c].width = w
    wq.freeze_panes = "A2"

    # About
    wa = wb.create_sheet("About")
    notes = ["About this report",
             "Generated by build_report.py from the FormWise database (made-up company, simulated data).",
             "Purchases and revenue come from billing data, not front-end events, because the purchase event has known gaps.",
             "Visitors and purchases are grouped by the month of the visitor's first visit; renewal revenue by payment month.",
             "Every KPI is a formula. Change 'Report month' on the KPI Summary sheet to update all figures.",
             "Data Quality flags months where purchase tracking disagrees with billing beyond the tolerance."]
    for i, n in enumerate(notes, 1):
        wa[f"A{i}"] = n
        wa[f"A{i}"].font = Font(name=FONT, bold=(i == 1), size=13 if i == 1 else 10)
    wa.column_dimensions["A"].width = 110
    for sh in wb.worksheets: set_fonts(sh)

    os.makedirs(out_dir, exist_ok=True)
    path = os.path.join(out_dir, f"FormWise_KPI_Report_{month}.xlsx")
    wb.save(path)

    # plain-language summary
    mon = base.groupby("month")[["visitors", "purchases", "new_revenue"]].sum()
    mon["conversion"] = mon.purchases / mon.visitors
    i = months.index(month); prev = months[i - 1] if i else None
    lines = [f"# FormWise KPI summary, {month}", ""]
    cur = mon.loc[month]
    lines.append(f"- Visitors: {cur.visitors:,.0f}; new subscriptions: {cur.purchases:,.0f}; "
                 f"conversion: {cur.conversion:.2%}.")
    if prev:
        p = mon.loc[prev]
        lines.append(f"- Versus {prev}: visitors {cur.visitors/p.visitors-1:+.1%}, purchases {cur.purchases/p.purchases-1:+.1%}, "
                     f"conversion {(cur.conversion-p.conversion)*100:+.2f} pts.")
    cd = base[base.month == month].groupby("channel")[["visitors", "purchases"]].sum()
    cd["conv"] = cd.purchases / cd.visitors
    if prev:
        pdv = base[base.month == prev].groupby("channel")[["visitors", "purchases"]].sum()
        cd["change_pts"] = (cd.conv - pdv.purchases / pdv.visitors) * 100
        mover = cd.change_pts.abs().idxmax()
        lines.append(f"- Biggest channel move: {mover} conversion {cd.loc[mover,'change_pts']:+.2f} pts "
                     f"(n={cd.loc[mover,'visitors']:,.0f} visitors, could be noise at this size).")
    flagged = dq[(dq.missing_user_id > 0) | ((dq.purchase_events / dq.billed_new_subs - 1).abs() > 0.03)]
    lines.append(f"- Data quality: {'CHECK ' + ', '.join(flagged.month) if len(flagged) else 'all months within tolerance'}.")
    md = os.path.join(out_dir, f"summary_{month}.md")
    open(md, "w").write("\n".join(lines) + "\n")
    con.close()
    return path, md


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--month", default="2026-06")
    ap.add_argument("--out", default=os.path.join(HERE, "output"))
    a = ap.parse_args()
    for p in build(a.month, a.out): print("wrote", p)
